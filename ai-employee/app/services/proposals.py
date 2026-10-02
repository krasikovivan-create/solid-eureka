"""Коммерческое предложение: структура, расчёт стоимости и DOCX-файл."""

from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from sqlalchemy import select

from app.database.models import Audit, CompanyProfile, Proposal
from app.database.session import SessionFactory
from app.services.tasks import NotFoundError
from app.services.timeutils import MONTHS

PROPOSAL_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "problem": {"type": "string"},
        "pain_points": {"type": "array", "items": {"type": "string"}},
        "solution": {"type": "string"},
        "benefits": {"type": "array", "items": {"type": "string"}},
        "stages": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "description": {"type": "string"},
                    "duration_weeks": {"type": "number"},
                    "effort_hours": {"type": "number"},
                },
                "required": ["name", "description", "duration_weeks", "effort_hours"],
                "additionalProperties": False,
            },
        },
        "next_steps": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["title", "problem", "pain_points", "solution", "benefits", "stages", "next_steps"],
    "additionalProperties": False,
}


def _num(value: Any) -> float:
    try:
        return max(0.0, float(value))
    except (TypeError, ValueError):
        return 0.0


def calculate(content: dict, hourly_rate: int) -> dict:
    """Считает стоимость этапов по ставке, общие сроки и итог."""
    stages = []
    for st in content.get("stages", []) or []:
        hours = round(_num(st.get("effort_hours")), 1)
        weeks = round(_num(st.get("duration_weeks")), 1)
        stages.append(
            {
                "name": str(st.get("name", "")),
                "description": str(st.get("description", "")),
                "duration_weeks": weeks,
                "effort_hours": hours,
                "cost": int(round(hours * hourly_rate, -2)),
            }
        )
    result = dict(content)
    result["stages"] = stages
    result["hourly_rate"] = hourly_rate
    result["total_weeks"] = round(sum(s["duration_weeks"] for s in stages), 1)
    result["total_hours"] = round(sum(s["effort_hours"] for s in stages), 1)
    result["total_cost"] = sum(s["cost"] for s in stages)
    return result


def money(value: int | float) -> str:
    return f"{int(value):,}".replace(",", " ") + " ₽"


def build_prompt(audit: Audit, company_context: str) -> str:
    report = audit.report or {}
    automations = "\n".join(
        f"- {a['process']}: {a['solution']} (сложность {a['complexity']}, "
        f"{a['duration_weeks']} нед., {a['effort_hours']} ч работы, "
        f"экономия {a['hours_saved_per_month']} ч/мес)"
        for a in report.get("automations", [])
    )
    processes = "\n".join(
        f"- {p['name']}: {p['description']} (~{p['hours_per_month']} ч/мес)"
        for p in report.get("processes", [])
    )
    return (
        f"{company_context}\n\n"
        f"Подготовь коммерческое предложение для клиента «{audit.client_name}» "
        "по итогам аудита.\n\n"
        f"Вывод аудита: {report.get('summary', '')}\n\n"
        f"Рутинные процессы:\n{processes}\n\n"
        f"Что автоматизировать:\n{automations}\n\n"
        f"Общая экономия: {report.get('total_hours_saved_per_month', 0)} ч/мес.\n\n"
        "Структура: title (заголовок КП), problem (проблема клиента, 2–4 предложения), "
        "pain_points (конкретные боли), solution (наше решение, 3–5 предложений), "
        "benefits (результаты для клиента с цифрами экономии), stages (этапы работ: название, "
        "что делаем, длительность в неделях, трудоёмкость в часах — стоимость посчитаем сами "
        "по ставке), next_steps (что делать клиенту дальше). Пиши по-русски, убедительно, "
        "без воды, в стиле нашей компании."
    )


def render_docx(content: dict, profile: CompanyProfile, client_name: str) -> bytes:
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    accent = RGBColor(0x1F, 0x4E, 0x79)

    def heading(text: str, level: int = 1) -> None:
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = accent

    today = date.today()
    company = profile.company_name or "Наша компания"
    p = doc.add_paragraph()
    run = p.add_run(company)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = accent
    if profile.company_description:
        doc.add_paragraph(profile.company_description).runs[0].italic = True

    title = doc.add_heading(content.get("title") or "Коммерческое предложение", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(
        f"Для: {client_name}\nДата: {today.day} {MONTHS[today.month - 1]} {today.year}"
    )

    heading("1. Проблема")
    doc.add_paragraph(content.get("problem", ""))
    for item in content.get("pain_points", []) or []:
        doc.add_paragraph(str(item), style="List Bullet")

    heading("2. Решение")
    doc.add_paragraph(content.get("solution", ""))
    if content.get("benefits"):
        doc.add_paragraph("Что получит клиент:").runs[0].bold = True
        for item in content["benefits"]:
            doc.add_paragraph(str(item), style="List Bullet")

    heading("3. Этапы, сроки и стоимость")
    stages = content.get("stages", []) or []
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["№", "Этап", "Срок, нед.", "Часы", "Стоимость"]
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for i, st in enumerate(stages, start=1):
        row = table.add_row().cells
        row[0].text = str(i)
        row[1].text = f"{st['name']}\n{st['description']}".strip()
        row[2].text = f"{st['duration_weeks']:g}"
        row[3].text = f"{st['effort_hours']:g}"
        row[4].text = money(st["cost"])
    total = table.add_row().cells
    total[1].text = "Итого"
    total[2].text = f"{content.get('total_weeks', 0):g}"
    total[3].text = f"{content.get('total_hours', 0):g}"
    total[4].text = money(content.get("total_cost", 0))
    for cell in total:
        for par in cell.paragraphs:
            for r in par.runs:
                r.bold = True

    heading("4. Сроки")
    doc.add_paragraph(
        f"Общая длительность проекта — около {content.get('total_weeks', 0):g} нед. "
        "Этапы могут частично идти параллельно, точный график согласуем на старте."
    )

    heading("5. Стоимость")
    doc.add_paragraph(
        f"Итоговая стоимость: {money(content.get('total_cost', 0))} "
        f"(трудоёмкость {content.get('total_hours', 0):g} ч по ставке "
        f"{money(content.get('hourly_rate', 0))}/ч). Оплата поэтапно."
    )

    if content.get("next_steps"):
        heading("6. Следующие шаги")
        for item in content["next_steps"]:
            doc.add_paragraph(str(item), style="List Number")

    contact = profile.agent_name
    if contact:
        doc.add_paragraph("")
        doc.add_paragraph(
            f"С уважением,\n{contact}{', ' + profile.agent_role if profile.agent_role else ''}\n"
            f"{company}"
        )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def proposal_filename(client_name: str) -> str:
    safe = re.sub(r"[^\w\- ]+", "", client_name, flags=re.UNICODE).strip().replace(" ", "_")
    return f"КП_{safe or 'клиент'}_{date.today():%Y-%m-%d}.docx"


def _ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


async def save_proposal(
    sf: SessionFactory,
    storage_dir: Path,
    audit: Audit,
    content: dict,
    data: bytes,
    filename: str,
) -> Proposal:
    _ensure_dir(storage_dir)
    async with sf() as s:
        proposal = Proposal(
            audit_id=audit.id,
            client_id=audit.client_id,
            title=str(content.get("title", ""))[:300],
            content=content,
            total_cost=int(content.get("total_cost", 0)),
        )
        s.add(proposal)
        await s.flush()
        path = storage_dir / f"{proposal.id}_{filename}"
        path.write_bytes(data)
        proposal.file_path = str(path)
        await s.commit()
        await s.refresh(proposal)
        return proposal


async def get_proposal(sf: SessionFactory, proposal_id: int) -> Proposal:
    async with sf() as s:
        proposal = await s.get(Proposal, proposal_id)
        if proposal is None:
            raise NotFoundError(f"КП #{proposal_id} не найдено")
        return proposal


async def list_proposals(sf: SessionFactory, limit: int = 20) -> list[Proposal]:
    async with sf() as s:
        rows = await s.scalars(select(Proposal).order_by(Proposal.id.desc()).limit(limit))
        return list(rows.all())

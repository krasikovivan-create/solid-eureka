"""Извлечение текста из PDF, DOCX и TXT."""

from __future__ import annotations

import io
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}
MAX_FILE_SIZE = 20 * 1024 * 1024  # лимит Telegram Bot API на скачивание


class UnsupportedDocumentError(ValueError):
    pass


def extension_of(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def is_supported(filename: str) -> bool:
    return extension_of(filename) in SUPPORTED_EXTENSIONS


def extract_text(filename: str, data: bytes) -> str:
    ext = extension_of(filename)
    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext in {".txt", ".md", ".csv"}:
        text = _decode_text(data)
    else:
        raise UnsupportedDocumentError(
            "Поддерживаются файлы PDF, DOCX, TXT (а также MD и CSV). "
            "Старый формат .doc сохраните как .docx."
        )
    text = _normalize(text)
    if not text.strip():
        raise UnsupportedDocumentError(
            "В файле не найден текст. Если это скан PDF — пришлите текстовую версию."
        )
    return text


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise UnsupportedDocumentError("PDF защищён паролем") from exc
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"[стр. {number}]\n{page_text}")
        return "\n\n".join(pages)
    except PdfReadError as exc:
        raise UnsupportedDocumentError("Не удалось прочитать PDF: файл повреждён") from exc


def _docx_text(data: bytes) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise UnsupportedDocumentError("Не удалось прочитать DOCX: файл повреждён") from exc
    except Exception as exc:  # zipfile.BadZipFile и подобные
        raise UnsupportedDocumentError("Не удалось прочитать DOCX: файл повреждён") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251", "koi8-r"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalize(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    result: list[str] = []
    blank = 0
    for line in lines:
        if line.strip():
            blank = 0
            result.append(line)
        else:
            blank += 1
            if blank <= 1:
                result.append("")
    return "\n".join(result).strip()


def chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    """Режет текст на фрагменты по абзацам, ~size символов с перекрытием."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        while len(para) > size:
            # Очень длинный абзац режем по предложениям/символам.
            cut = para.rfind(". ", 0, size)
            cut = cut + 1 if cut > size // 2 else size
            piece, para = para[:cut].strip(), para[cut:].strip()
            if current:
                chunks.append(current)
                current = ""
            chunks.append(piece)
        if len(current) + len(para) + 2 <= size:
            current = f"{current}\n\n{para}" if current else para
        else:
            if current:
                chunks.append(current)
            tail = current[-overlap:] if current and overlap else ""
            current = f"{tail}\n\n{para}".strip() if tail else para
    if current:
        chunks.append(current)
    return chunks

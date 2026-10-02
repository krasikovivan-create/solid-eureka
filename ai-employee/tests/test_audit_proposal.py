import io
import json

import docx
import pytest

from app.agent import workflows
from app.agent.llm import LLMError
from app.services import audit as audit_service
from app.services import clients as client_service
from app.services import proposals as proposal_service
from app.tools import OutAuditStarted, OutFile, registry
from tests.conftest import OWNER_ID, response, text_block

REPORT = {
    "summary": "Небольшая сеть салонов, много ручной записи клиентов.",
    "processes": [
        {
            "name": "Запись клиентов",
            "description": "Администратор отвечает в WhatsApp",
            "who": "администратор",
            "hours_per_month": 80,
        },
        {
            "name": "Отчёты",
            "description": "Ручная сводка в Excel",
            "who": "управляющий",
            "hours_per_month": 12,
        },
    ],
    "automations": [
        {
            "process": "Запись клиентов",
            "solution": "ИИ-бот записи в WhatsApp и Telegram",
            "complexity": "средняя",
            "duration_weeks": 3,
            "effort_hours": 60,
            "hours_saved_per_month": 60,
        },
        {
            "process": "Отчёты",
            "solution": "Автоматическая сводка из CRM",
            "complexity": "низкая",
            "duration_weeks": 1,
            "effort_hours": 16,
            "hours_saved_per_month": 10,
        },
    ],
    "quick_wins": ["Шаблоны ответов"],
    "risks": ["Нужен доступ к CRM"],
}

PROPOSAL = {
    "title": "Автоматизация записи клиентов",
    "problem": "Администраторы тратят 80 часов в месяц на ручную запись.",
    "pain_points": ["Потерянные заявки", "Ошибки в записи"],
    "solution": "Внедряем ИИ-бота записи и автоматические отчёты.",
    "benefits": ["Экономия 70 часов в месяц"],
    "stages": [
        {
            "name": "Анализ",
            "description": "Сбор требований",
            "duration_weeks": 1,
            "effort_hours": 10,
        },
        {
            "name": "Разработка бота",
            "description": "Бот записи",
            "duration_weeks": 3,
            "effort_hours": 60,
        },
    ],
    "next_steps": ["Подписать договор", "Дать доступ к CRM"],
}


async def answer_all(sf, audit_id):
    audit = None
    for i in range(len(audit_service.AUDIT_QUESTIONS)):
        audit = await audit_service.answer(sf, audit_id, f"ответ {i}" if i != 3 else None)
    return audit


async def test_audit_interview_state_in_db(onboarded):
    sf = onboarded.sf
    audit = await audit_service.start_audit(sf, OWNER_ID, "Салон Красоты")
    assert audit_service.current_question(audit) == audit_service.AUDIT_QUESTIONS[0]
    audit = await audit_service.answer(sf, audit.id, "Сеть из 3 салонов")
    # состояние читается из БД заново (как после перезапуска)
    active = await audit_service.get_active_audit(sf, OWNER_ID)
    assert active.id == audit.id and active.step == 1
    assert active.answers[0]["answer"] == "Сеть из 3 салонов"
    audit = await answer_all(sf, audit.id)
    assert audit_service.is_interview_finished(audit)
    assert audit.answers[4]["answer"] == "—"  # пропущенный вопрос (индекс 3 + 1 первый ответ)


async def test_new_audit_cancels_previous(onboarded):
    sf = onboarded.sf
    a1 = await audit_service.start_audit(sf, OWNER_ID, "A")
    a2 = await audit_service.start_audit(sf, OWNER_ID, "B")
    assert (await audit_service.get_audit(sf, a1.id)).status == "cancelled"
    assert (await audit_service.get_active_audit(sf, OWNER_ID)).id == a2.id
    await audit_service.cancel_audit(sf, a2.id)
    assert await audit_service.get_active_audit(sf, OWNER_ID) is None


async def test_audit_report_generation(onboarded, fake_llm):
    sf = onboarded.sf
    client = await client_service.create_client(sf, "Салон Красоты")
    audit = await audit_service.start_audit(sf, OWNER_ID, client.company, client.id)
    await answer_all(sf, audit.id)
    fake_llm.add(
        response(text_block(json.dumps(REPORT, ensure_ascii=False)), model="claude-sonnet-5-5")
    )
    audit = await workflows.generate_audit_report(onboarded, audit.id, OWNER_ID)
    call = fake_llm.calls[-1]
    assert call["model"] == "claude-sonnet-5-5"
    assert call["output_config"]["format"]["type"] == "json_schema"
    assert "ответ 0" in call["messages"][0]["content"]
    assert audit.status == "completed"
    assert audit.report["total_hours_saved_per_month"] == 70
    assert audit.report["total_routine_hours_per_month"] == 92
    text = audit_service.format_report(audit)
    assert "Карта рутинных процессов" in text and "70 ч в месяц" in text
    assert "средняя" in text and "3 нед." in text
    # карточка клиента обновлена
    client = await client_service.get_client(sf, client.id)
    assert client.status == "audit"
    assert any("аудит" in i.summary.lower() for i in client.interactions)


async def test_report_normalization_handles_garbage():
    r = audit_service.normalize_report(
        {
            "automations": [
                {"process": "x", "complexity": "космическая", "hours_saved_per_month": "abc"}
            ]
        }
    )
    assert r["automations"][0]["complexity"] == "средняя"
    assert r["automations"][0]["hours_saved_per_month"] == 0
    assert r["processes"] == []


async def test_proposal_calculation_and_docx(onboarded):
    from app.services.profile import get_profile

    profile = await get_profile(onboarded.sf)
    content = proposal_service.calculate(PROPOSAL, 3000)
    assert content["stages"][0]["cost"] == 30000
    assert content["total_cost"] == 210000
    assert content["total_weeks"] == 4 and content["total_hours"] == 70
    data = proposal_service.render_docx(content, profile, "Салон Красоты")
    doc = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    for part in ("1. Проблема", "2. Решение", "3. Этапы", "4. Сроки", "5. Стоимость", "ИИ Лаб"):
        assert part in text
    assert "210 000 ₽" in text
    table_text = "\n".join(c.text for row in doc.tables[0].rows for c in row.cells)
    assert "Разработка бота" in table_text and "180 000 ₽" in table_text


async def test_proposal_workflow_and_tool(tool_ctx_factory, fake_llm):
    ctx = tool_ctx_factory()
    sf = ctx.app.sf
    client = await client_service.create_client(sf, "Салон Красоты")
    audit = await audit_service.start_audit(sf, OWNER_ID, client.company, client.id)
    # КП по незавершённому аудиту — понятная ошибка
    out, err = await registry.execute(ctx, "create_proposal", {"audit_id": audit.id})
    assert err and "не завершён" in out
    await answer_all(sf, audit.id)
    await audit_service.save_report(sf, audit.id, REPORT)
    fake_llm.add(response(text_block(json.dumps(PROPOSAL, ensure_ascii=False))))
    out, err = await registry.execute(ctx, "create_proposal", {"audit_id": audit.id})
    assert not err, out
    data = json.loads(out)
    assert data["total_cost_rub"] == 210000
    file = ctx.outbox[-1]
    assert isinstance(file, OutFile) and file.filename.endswith(".docx")
    assert file.filename.startswith("КП_Салон_Красоты")
    docx.Document(io.BytesIO(file.data))  # файл открывается
    proposals = await proposal_service.list_proposals(sf)
    assert proposals[0].total_cost == 210000
    client = await client_service.get_client(sf, client.id)
    assert client.status == "proposal"


async def test_start_audit_tool(tool_ctx_factory):
    ctx = tool_ctx_factory()
    await client_service.create_client(ctx.app.sf, "Ромашка")
    out, err = await registry.execute(ctx, "start_audit", {"client_name": "ромашка"})
    assert not err
    data = json.loads(out)
    assert data["client"] == "Ромашка"
    assert isinstance(ctx.outbox[-1], OutAuditStarted)
    audit = await audit_service.get_active_audit(ctx.app.sf, OWNER_ID)
    assert audit.client_id is not None
    out, err = await registry.execute(ctx, "start_audit", {"client_name": "Новый клиент"})
    assert not err and json.loads(out)["client"] == "Новый клиент"
    out, _ = await registry.execute(ctx, "list_audits", {})
    assert len(json.loads(out)) == 2
    out, err = await registry.execute(ctx, "start_audit", {})
    assert err


async def test_get_audit_report_tool(tool_ctx_factory):
    ctx = tool_ctx_factory()
    audit = await audit_service.start_audit(ctx.app.sf, OWNER_ID, "X")
    out, _ = await registry.execute(ctx, "get_audit_report", {"audit_id": audit.id})
    assert json.loads(out)["report"] is None
    await audit_service.save_report(ctx.app.sf, audit.id, REPORT)
    out, _ = await registry.execute(ctx, "get_audit_report", {"audit_id": audit.id})
    assert json.loads(out)["report"]["total_hours_saved_per_month"] == 70


async def test_bad_json_from_model(onboarded, fake_llm):
    audit = await audit_service.start_audit(onboarded.sf, OWNER_ID, "X")
    fake_llm.add(response(text_block("не json вовсе")))
    with pytest.raises(LLMError):
        await workflows.generate_audit_report(onboarded, audit.id, OWNER_ID)

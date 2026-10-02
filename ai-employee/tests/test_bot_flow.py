"""Сквозные сценарии через диспетчер aiogram с поддельным Telegram."""

import io
import itertools
import json
from datetime import UTC, datetime

import docx
import pytest
from aiogram.methods import EditMessageText, SendDocument
from aiogram.types import CallbackQuery, Chat, Document, Message, Update, User

from app.agent.agent import Agent
from app.bot import build_dispatcher
from app.services import audit as audit_service
from app.services import clients as client_service
from app.services import profile as profile_service
from app.services import tasks as task_service
from tests.conftest import OWNER_ID, response, text_block, tool_block
from tests.test_audit_proposal import PROPOSAL, REPORT

_ids = itertools.count(1)


def user(uid: int) -> User:
    return User(id=uid, is_bot=False, first_name="Иван", username=f"u{uid}")


def message(text: str | None = None, uid: int = OWNER_ID, **extra) -> Message:
    return Message(
        message_id=next(_ids),
        date=datetime.now(UTC),
        chat=Chat(id=uid, type="private"),
        from_user=user(uid),
        text=text,
        **extra,
    )


def msg_update(text: str, uid: int = OWNER_ID) -> Update:
    return Update(update_id=next(_ids), message=message(text, uid))


def cb_update(data: str, uid: int = OWNER_ID) -> Update:
    return Update(
        update_id=next(_ids),
        callback_query=CallbackQuery(
            id=str(next(_ids)),
            from_user=user(uid),
            chat_instance="ci",
            data=data,
            message=message("старое сообщение", uid),
        ),
    )


@pytest.fixture
def dp(app):
    return build_dispatcher(app, Agent(app))


@pytest.fixture
def dp_ready(onboarded):
    return build_dispatcher(onboarded, Agent(onboarded))


async def test_first_start_onboarding_full(dp, bot, session, app):
    await dp.feed_update(bot, msg_update("/start"))
    assert "Шаг 1 из 5" in session.sent_texts()[-1]
    profile = await profile_service.get_profile(app.sf)
    assert profile.owner_id == OWNER_ID  # первый нажавший /start — владелец

    for answer in ("ИИ Лаб", "Внедряем ИИ в бизнес", "Анна", "ассистент руководителя"):
        await dp.feed_update(bot, msg_update(answer))
    last = session.sent_messages()[-1]
    assert "Шаг 5 из 5" in last.text
    assert last.reply_markup is not None  # кнопки стилей

    await dp.feed_update(bot, cb_update("style:business"))
    done = session.sent_messages()[-1]
    assert "Готово" in done.text and "09:00" in done.text
    profile = await profile_service.get_profile(app.sf)
    assert profile.onboarded
    assert profile.company_name == "ИИ Лаб"
    assert profile.company_description == "Внедряем ИИ в бизнес"
    assert profile.agent_name == "Анна" and profile.agent_role == "ассистент руководителя"
    assert profile.communication_style.startswith("Деловой")
    assert app.scheduler.scheduler.get_job("daily:morning") is not None

    # повторный /start — главное меню
    await dp.feed_update(bot, msg_update("/start"))
    assert "Анна" in session.sent_texts()[-1]


async def test_onboarding_survives_restart(dp, bot, session, app):
    await dp.feed_update(bot, msg_update("/start"))
    await dp.feed_update(bot, msg_update("ИИ Лаб"))
    # «перезапуск»: новый диспетчер, шаг берётся из БД
    dp2 = build_dispatcher(app, Agent(app))
    await dp2.feed_update(bot, msg_update("/start"))
    assert "Шаг 2 из 5" in session.sent_texts()[-1]


async def test_strangers_are_blocked(dp, bot, session, app):
    await dp.feed_update(bot, msg_update("/start", uid=OWNER_ID))
    session.clear()
    await dp.feed_update(bot, msg_update("/start", uid=666))
    assert "нет доступа" in session.sent_texts()[-1]
    assert "666" in session.sent_texts()[-1]
    profile = await profile_service.get_profile(app.sf)
    assert profile.owner_id == OWNER_ID


async def test_allowed_user_ids_whitelist(app, bot, session):
    app.settings.allowed_user_ids = [555]
    dp = build_dispatcher(app, Agent(app))
    await dp.feed_update(bot, msg_update("/start", uid=OWNER_ID))
    assert "нет доступа" in session.sent_texts()[-1]
    await dp.feed_update(bot, msg_update("/start", uid=555))
    assert "Шаг 1 из 5" in session.sent_texts()[-1]
    assert (await profile_service.get_profile(app.sf)).owner_id is None


async def test_free_text_goes_to_agent_and_tools(dp_ready, bot, session, onboarded, fake_llm):
    fake_llm.add(
        response(
            tool_block(
                "create_task", {"title": "Позвонить Иванову", "remind_at": "2030-05-05T10:00"}
            )
        ),
        response(text_block("**Готово!** Напомню 5 мая в 10:00.")),
    )
    await dp_ready.feed_update(bot, msg_update("напомни 5 мая 2030 в 10 позвонить Иванову"))
    assert "<b>Готово!</b>" in session.sent_texts()[-1]
    tasks = await task_service.list_tasks(onboarded.sf, OWNER_ID, "Europe/Moscow")
    assert tasks[0].title == "Позвонить Иванову"


async def test_llm_error_message_is_shown(dp_ready, bot, session, onboarded):
    onboarded.settings.daily_budget_usd = 0.000001
    from app.services.usage import record_usage

    await record_usage(onboarded.sf, "claude-sonnet-5-5", 10_000, 1_000)
    await dp_ready.feed_update(bot, msg_update("привет"))
    assert "лимит" in session.sent_texts()[-1]


async def test_unexpected_error_is_reported_not_silent(dp_ready, bot, session, fake_llm):
    fake_llm.add(RuntimeError("boom"))
    await dp_ready.feed_update(bot, msg_update("привет"))
    assert "Что-то пошло не так" in session.sent_texts()[-1]


async def test_delete_with_confirmation_buttons(dp_ready, bot, session, onboarded, fake_llm):
    task = await task_service.create_task(onboarded.sf, OWNER_ID, "Лишняя задача")
    fake_llm.add(
        response(tool_block("delete_task", {"task_id": task.id})),
        response(text_block("Подтвердите удаление.")),
    )
    await dp_ready.feed_update(bot, msg_update("удали задачу про лишнее"))
    confirm = session.sent_messages()[-1]
    assert "Удалить задачу" in confirm.text
    buttons = [b.callback_data for row in confirm.reply_markup.inline_keyboard for b in row]
    assert f"delok:task:{task.id}" in buttons and "delno" in buttons
    # отмена
    await dp_ready.feed_update(bot, cb_update("delno"))
    assert await task_service.get_task(onboarded.sf, task.id)
    # подтверждение
    await dp_ready.feed_update(bot, cb_update(f"delok:task:{task.id}"))
    edits = [r for r in session.requests if isinstance(r, EditMessageText)]
    assert "удалена" in edits[-1].text
    with pytest.raises(task_service.NotFoundError):
        await task_service.get_task(onboarded.sf, task.id)


async def test_delete_client_from_card(dp_ready, bot, session, onboarded):
    client = await client_service.create_client(onboarded.sf, "Ромашка")
    await dp_ready.feed_update(bot, cb_update(f"client:view:{client.id}"))
    await dp_ready.feed_update(bot, cb_update(f"del:client:{client.id}"))
    assert "Удалить клиента" in session.sent_texts()[-1]
    await dp_ready.feed_update(bot, cb_update(f"delok:client:{client.id}"))
    assert await client_service.count_clients(onboarded.sf) == 0


async def test_main_menu_sections(dp_ready, bot, session, onboarded):
    await client_service.create_client(onboarded.sf, "Ромашка")
    for section in (
        "main", "tasks", "today", "clients", "pipeline", "audit", "kb", "texts",
        "memory", "usage", "settings", "help",
    ):  # fmt: skip
        session.clear()
        await dp_ready.feed_update(bot, cb_update(f"menu:{section}"))
        edits = [r for r in session.requests if isinstance(r, EditMessageText)]
        assert edits, section
    for cmd in (
        "/menu",
        "/help",
        "/tasks",
        "/today",
        "/clients",
        "/kb",
        "/usage",
        "/settings",
        "/audit",
    ):
        session.clear()
        await dp_ready.feed_update(bot, msg_update(cmd))
        assert session.sent_texts(), cmd


async def test_help_has_examples(dp_ready, bot, session):
    await dp_ready.feed_update(bot, msg_update("/help"))
    text = session.sent_texts()[-1]
    assert "Напомни завтра в 10 позвонить Иванову" in text


async def test_settings_change_time(dp_ready, bot, session, onboarded):
    await dp_ready.feed_update(bot, cb_update("set:morning_time"))
    await dp_ready.feed_update(bot, msg_update("25:99"))
    assert "ЧЧ:ММ" in session.sent_texts()[-1]
    await dp_ready.feed_update(bot, msg_update("8:45"))
    assert "Сохранено" in session.sent_texts()[-1]
    profile = await profile_service.get_profile(onboarded.sf)
    assert profile.morning_time == "08:45"
    job = onboarded.scheduler.scheduler.get_job("daily:morning")
    fields = {f.name: str(f) for f in job.trigger.fields}
    assert fields["hour"] == "8" and fields["minute"] == "45"
    # переключатель
    await dp_ready.feed_update(bot, cb_update("toggle:evening_enabled"))
    assert not (await profile_service.get_profile(onboarded.sf)).evening_enabled
    # ставка
    await dp_ready.feed_update(bot, cb_update("set:hourly_rate"))
    await dp_ready.feed_update(bot, msg_update("4 000 ₽"))
    assert (await profile_service.get_profile(onboarded.sf)).hourly_rate == 4000


async def test_reminder_buttons_done_and_snooze(dp_ready, bot, session, onboarded):
    task = await task_service.create_task(onboarded.sf, OWNER_ID, "Позвонить")
    await dp_ready.feed_update(bot, cb_update(f"snooze:{task.id}:60"))
    task = await task_service.get_task(onboarded.sf, task.id)
    assert task.remind_at is not None and not task.reminder_sent
    assert "Напомню" in session.sent_texts()[-1]
    await dp_ready.feed_update(bot, cb_update(f"snooze:{task.id}:tomorrow"))
    await dp_ready.feed_update(bot, cb_update(f"task:done:{task.id}"))
    assert (await task_service.get_task(onboarded.sf, task.id)).status == "done"


async def test_document_upload_and_question(dp_ready, bot, session, onboarded, fake_llm):
    session.files["kb1"] = "Оплата: 50% предоплата, остаток после сдачи.".encode()
    upd = Update(
        update_id=next(_ids),
        message=message(
            None,
            document=Document(
                file_id="kb1", file_unique_id="u1", file_name="Условия.txt", file_size=60
            ),
        ),
    )
    await dp_ready.feed_update(bot, upd)
    assert "добавлен в базу знаний" in session.sent_texts()[-1]
    docs = await onboarded.kb.list_documents()
    assert docs[0][0].filename == "Условия.txt"

    def check_tool_result(kwargs):
        result = kwargs["messages"][-1]["content"][0]["content"]
        assert "предоплата" in result and "Условия.txt" in result
        return response(text_block("50% предоплата (📄 Условия.txt)"))

    fake_llm.add(response(tool_block("search_knowledge", {"query": "оплата"})), check_tool_result)
    await dp_ready.feed_update(bot, msg_update("Какие условия оплаты?"))
    assert "📄 Условия.txt" in session.sent_texts()[-1]


async def test_unsupported_document(dp_ready, bot, session):
    upd = Update(
        update_id=next(_ids),
        message=message(
            None, document=Document(file_id="x", file_unique_id="x", file_name="photo.png")
        ),
    )
    await dp_ready.feed_update(bot, upd)
    assert "не читаю" in session.sent_texts()[-1]


async def test_audit_interview_then_proposal_docx(dp_ready, bot, session, onboarded, fake_llm):
    await client_service.create_client(onboarded.sf, "Салон Красоты")
    await dp_ready.feed_update(bot, cb_update("audit:new"))
    await dp_ready.feed_update(bot, msg_update("Салон Красоты"))
    assert "Вопрос 1 из 10" in session.sent_texts()[-1]
    audit = await audit_service.get_active_audit(onboarded.sf, OWNER_ID)
    assert audit.client_id is not None

    await dp_ready.feed_update(bot, msg_update("Сеть салонов, 20 сотрудников"))
    assert "Вопрос 2 из 10" in session.sent_texts()[-1]
    await dp_ready.feed_update(bot, cb_update("audit:skip"))
    assert "Вопрос 3 из 10" in session.sent_texts()[-1]
    fake_llm.add(response(text_block(json.dumps(REPORT, ensure_ascii=False))))
    for i in range(3, 11):
        await dp_ready.feed_update(bot, msg_update(f"Ответ {i}"))
    report_msg = session.sent_messages()[-1]
    assert "Аудит процессов" in report_msg.text and "70 ч в месяц" in report_msg.text
    buttons = [b.callback_data for row in report_msg.reply_markup.inline_keyboard for b in row]
    assert f"audit:proposal:{audit.id}" in buttons

    fake_llm.add(response(text_block(json.dumps(PROPOSAL, ensure_ascii=False))))
    await dp_ready.feed_update(bot, cb_update(f"audit:proposal:{audit.id}"))
    docs = [r for r in session.requests if isinstance(r, SendDocument)]
    assert docs, "КП не отправлено"
    file = docs[-1].document
    assert file.filename.endswith(".docx")
    parsed = docx.Document(io.BytesIO(file.data))
    assert any("Проблема" in p.text for p in parsed.paragraphs)
    assert "210 000 ₽" in docs[-1].caption


async def test_audit_stop_and_cancel(dp_ready, bot, session, onboarded):
    await dp_ready.feed_update(bot, cb_update("audit:new"))
    await dp_ready.feed_update(bot, msg_update("Клиент"))
    await dp_ready.feed_update(bot, cb_update("audit:stop"))
    assert await audit_service.get_active_audit(onboarded.sf, OWNER_ID) is None
    await audit_service.start_audit(onboarded.sf, OWNER_ID, "Клиент2")
    await dp_ready.feed_update(bot, msg_update("/cancel"))
    assert await audit_service.get_active_audit(onboarded.sf, OWNER_ID) is None


async def test_agent_starts_audit_and_bot_asks_first_question(
    dp_ready, bot, session, onboarded, fake_llm
):
    fake_llm.add(
        response(tool_block("start_audit", {"client_name": "Ромашка"})),
        response(text_block("Начинаем аудит Ромашки.")),
    )
    await dp_ready.feed_update(bot, msg_update("Проведи аудит процессов для Ромашки"))
    texts = session.sent_texts()
    assert "Начинаем аудит Ромашки." in texts[-2]
    assert "Вопрос 1 из 10" in texts[-1]


async def test_write_text_reply_delivered(dp_ready, bot, session, fake_llm):
    fake_llm.add(
        response(tool_block("write_text", {"kind": "post", "brief": "пост про ИИ"})),
        response(text_block("Пост выше. Нужна правка?")),
        # порядок вызовов: агент → write_text (smart) → агент
    )
    fake_llm.script.insert(1, response(text_block("ИИ экономит время! #ИИ")))
    await dp_ready.feed_update(bot, msg_update("Напиши пост про ИИ"))
    texts = session.sent_texts()
    assert "Пост выше" in texts[-2] and "ИИ экономит время" in texts[-1]


async def test_new_command_clears_history(dp_ready, bot, session, onboarded):
    from app.services import memory as memory_service

    await memory_service.add_message(onboarded.sf, OWNER_ID, "user", "hi")
    await dp_ready.feed_update(bot, msg_update("/new"))
    assert await memory_service.unsummarized_messages(onboarded.sf, OWNER_ID) == []
